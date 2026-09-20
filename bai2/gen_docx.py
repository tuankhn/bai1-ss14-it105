#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate bai_tap_nhac_uong_thuoc.docx for RikkeiCare medicine reminder exercise."""
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def wireframe_box(doc, title, lines, width_cm=15):
    table = doc.add_table(rows=1 + len(lines), cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    title_cell = table.rows[0].cells[0]
    title_cell.text = title
    title_cell.paragraphs[0].runs[0].bold = True
    set_cell_shading(title_cell, 'D9D9D9')
    for i, line in enumerate(lines):
        cell = table.rows[i + 1].cells[0]
        cell.text = line
    for row in table.rows:
        row.cells[0].width = Cm(width_cm)
    doc.add_paragraph()


doc = Document()

title = doc.add_heading('BÀI TẬP: THIẾT KẾ TÍNH NĂNG NHẮC UỐNG THUỐC RIKKEICARE', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Họ và tên: .....................................       Lớp: .....................')
doc.add_paragraph()

# Bước 1
add_heading(doc, 'Bước 1 — Phân tích Use Case & Xác định yếu tố UI', level=1)
doc.add_paragraph(
    'Use Case tóm tắt: Đến giờ uống thuốc đã hẹn, hệ thống tự hiện thông báo nhắc nhở kèm tên thuốc → '
    'Bệnh nhân xem thông báo, nhấn "Đã uống" → Hệ thống ghi nhận và đóng thông báo.'
)
p = doc.add_paragraph()
p.add_run('Actor: ').bold = True
p.add_run('Bệnh nhân mãn tính, đa phần lớn tuổi (người dùng dễ bối rối với giao diện rườm rà, chữ nhỏ).')

doc.add_paragraph('Bảng xác định thành phần UI:')

table1 = doc.add_table(rows=3, cols=3)
table1.style = 'Table Grid'
hdr = table1.rows[0].cells
hdr[0].text = 'Yếu tố trong Use Case'
hdr[1].text = 'Thành phần UI tương ứng'
hdr[2].text = 'Ghi chú thiết kế (Clarity & Simplicity)'
for c in hdr:
    c.paragraphs[0].runs[0].bold = True
    set_cell_shading(c, 'D9D9D9')

rows_data = [
    ('Khối hiển thị tên thuốc', 'Text/Label lớn (icon viên thuốc + tên thuốc + liều dùng), đặt giữa khối thông báo (Modal/Card)', 'Cỡ chữ lớn (>=20px), độ tương phản cao, không kèm thông tin thừa'),
    ('Nút phản hồi "Đã uống"', 'Nút bấm (Button) to, chiếm toàn bộ chiều ngang, chỉ 1 nút duy nhất', 'Chỉ 1 lựa chọn duy nhất — giảm thời gian ra quyết định (Định luật Hick)'),
]
for i, (a, b, c) in enumerate(rows_data, start=1):
    cells = table1.rows[i].cells
    cells[0].text = a
    cells[1].text = b
    cells[2].text = c

doc.add_paragraph()
add_heading(doc, 'Bước 2 — Wireframe theo luồng', level=1)
doc.add_paragraph('Khung nền: Màn hình chính RikkeiCare (Xin chào, Bác Nguyễn Văn A) được bổ sung '
                   'khối thông báo Nhắc uống thuốc hiện lên dạng Modal/Card nổi giữa màn hình.')

add_heading(doc, 'Khung 1 — Thông báo nhắc thuốc vừa hiện lên', level=2)
wireframe_box(doc, 'RikkeiCare — Màn hình chính', [
    'Xin chào, Bác Nguyễn Văn A',
    '[ Lịch khám sắp tới: 08:00 - 15/09/2026 - BS. Trần B ]',
    '',
    '===== THÔNG BÁO NHẮC THUỐC (Modal, chữ to, nổi bật) =====',
    '   🔔 ĐÃ ĐẾN GIỜ UỐNG THUỐC',
    '   💊 Thuốc huyết áp Amlodipine 5mg',
    '',
    '   [        ĐÃ UỐNG        ]   (nút to, chiếm toàn chiều ngang)',
    '=========================================================',
])

add_heading(doc, 'Khung 2 — Sau khi nhấn "Đã uống"', level=2)
wireframe_box(doc, 'RikkeiCare — Màn hình chính', [
    'Xin chào, Bác Nguyễn Văn A',
    '[ Lịch khám sắp tới: 08:00 - 15/09/2026 - BS. Trần B ]',
    '',
    'Thông báo nhắc thuốc đã đóng lại.',
    '✔ Xác nhận (banner xanh lá, chữ to, hiện 2-3 giây):',
    '   "Đã ghi nhận: Bạn đã uống Amlodipine 5mg lúc 08:05."',
])

doc.add_paragraph()
add_heading(doc, 'Bước 3 — Tinh chỉnh theo nguyên tắc UI (Định luật Hick / Simplicity)', level=1)
p3 = doc.add_paragraph()
p3.add_run(
    'Nếu thêm 4-5 nút lựa chọn phản hồi (ví dụ: Đã uống, Uống trễ, Bỏ qua, Nhắc lại sau, Không có thuốc) '
    'thay vì chỉ 1 nút "Đã uống", theo Định luật Hick, thời gian ra quyết định sẽ tăng lên vì bệnh nhân '
    'lớn tuổi phải đọc và cân nhắc nhiều lựa chọn cùng lúc, dễ gây bối rối, chọn nhầm hoặc bỏ qua thông '
    'báo hoàn toàn — vi phạm nguyên tắc Simplicity và làm giảm tính Usability (dễ dùng) của tính năng.'
)

doc.save('bai_tap_nhac_uong_thuoc.docx')
print('Done')
