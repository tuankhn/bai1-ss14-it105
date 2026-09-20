#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate bai_tap_danh_gia_san_pham.docx for RikkeiShop review feature exercise."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def wireframe_box(doc, title, lines, width_cm=15, dashed=True):
    table = doc.add_table(rows=1 + len(lines), cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    title_cell = table.rows[0].cells[0]
    title_cell.text = title
    title_cell.paragraphs[0].runs[0].bold = True
    set_cell_shading(title_cell, 'D9D9D9')
    for i, line in enumerate(lines):
        cell = table.rows[i + 1].cells[0]
        cell.text = line
    for row in table.rows:
        row.cells[0].width = Cm(width_cm)
    doc.add_paragraph()


doc = Document()

# Title
title = doc.add_heading('BÀI TẬP: THIẾT KẾ TÍNH NĂNG ĐÁNH GIÁ SẢN PHẨM RIKKEISHOP', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Họ và tên: .....................................       Lớp: .....................')
doc.add_paragraph()

# Bước 1
add_heading(doc, 'Bước 1 — Phân tích Use Case & Xác định yếu tố UI', level=1)
doc.add_paragraph(
    'Use Case tóm tắt: Khách hàng đã mua sản phẩm nhấn nút "Viết đánh giá" → Hệ thống hiện form '
    'gồm chọn số sao (1-5) và ô nhập tiêu đề ngắn → Khách hàng điền xong, nhấn "Gửi đánh giá" → '
    'Hệ thống báo thành công.'
)
p = doc.add_paragraph()
p.add_run('Actor: ').bold = True
p.add_run('Khách hàng đã mua sản phẩm (Customer đã đặt hàng thành công).')

doc.add_paragraph('Bảng ánh xạ hành động nhập liệu sang thành phần UI:')

table1 = doc.add_table(rows=3, cols=3)
table1.style = 'Table Grid'
hdr = table1.rows[0].cells
hdr[0].text = 'Hành động nhập liệu (Use Case)'
hdr[1].text = 'Thành phần UI tương ứng'
hdr[2].text = 'Ghi chú thiết kế'
for c in hdr:
    c.paragraphs[0].runs[0].bold = True
    set_cell_shading(c, 'D9D9D9')

rows_data = [
    ('Chọn số sao đánh giá (1-5)', 'Star Rating component (5 icon ngôi sao, bấm chọn)', 'Bắt buộc chọn tối thiểu 1 sao trước khi cho phép gửi'),
    ('Nhập tiêu đề ngắn', 'Text Input (ô nhập văn bản 1 dòng) có placeholder gợi ý', 'Giới hạn độ dài, hiển thị số ký tự còn lại'),
]
for i, (a, b, c) in enumerate(rows_data, start=1):
    cells = table1.rows[i].cells
    cells[0].text = a
    cells[1].text = b
    cells[2].text = c

doc.add_paragraph()
add_heading(doc, 'Bước 2 — Wireframe theo luồng', level=1)
doc.add_paragraph('Khung nền: Trang Chi tiết sản phẩm RikkeiShop (Áo khoác nam Rikkei Basic — 459.000đ) '
                   'được bổ sung nút "Viết đánh giá" ngay dưới nút "Thêm vào giỏ".')

add_heading(doc, 'Khung 1 — Form "Viết đánh giá" đang mở', level=2)
wireframe_box(doc, 'RikkeiShop — Chi tiết sản phẩm', [
    'Áo khoác nam Rikkei Basic — 459.000đ',
    '[ Ảnh sản phẩm / Mô tả chi tiết ]',
    '[ Thêm vào giỏ ]     [ Viết đánh giá ] (đã bấm, đang active)',
    '--- Form đánh giá ---',
    'Chọn số sao: ☆ ☆ ☆ ☆ ☆   (bấm để tô sao, ví dụ đã chọn 4/5: ★ ★ ★ ★ ☆)',
    'Tiêu đề đánh giá: [ Nhập tiêu đề ngắn.................... ]  (0/50 ký tự)',
    '[ Hủy ]                              [ Gửi đánh giá ]',
])

add_heading(doc, 'Khung 2 — Sau khi nhấn "Gửi đánh giá"', level=2)
wireframe_box(doc, 'RikkeiShop — Chi tiết sản phẩm', [
    'Áo khoác nam Rikkei Basic — 459.000đ',
    '[ Ảnh sản phẩm / Mô tả chi tiết ]',
    '[ Thêm vào giỏ ]     [ Viết đánh giá ]',
    '✔ Thông báo (banner xanh lá, góc trên hoặc trong khu vực đánh giá):',
    '   "Cảm ơn bạn! Đánh giá của bạn đã được gửi thành công."',
    'Danh sách đánh giá được cập nhật, hiển thị đánh giá vừa gửi:',
    '  ★ ★ ★ ★ ☆  "Áo đẹp, form chuẩn"  — Bạn (vừa xong)',
])

doc.add_paragraph()
add_heading(doc, 'Bước 3 — Tinh chỉnh theo nguyên tắc UI (Feedback)', level=1)
p3 = doc.add_paragraph()
p3.add_run(
    'Nếu ở Khung 2 hệ thống chỉ đóng form lại mà không có bất kỳ thông báo nào, đây là trường hợp '
    'Good UI nhưng Bad UX: giao diện có thể vẫn gọn gàng, đẹp mắt (Good UI), nhưng người dùng không '
    'nhận được phản hồi (Feedback) rõ ràng nên không biết đánh giá đã được gửi thành công hay chưa, '
    'dẫn đến lo lắng, gửi lại nhiều lần hoặc mất niềm tin vào hệ thống — vi phạm nguyên tắc Feedback '
    'và giảm tính Utility (hữu ích) của tính năng.'
)

doc.save('bai_tap_danh_gia_san_pham.docx')
print('Done')
